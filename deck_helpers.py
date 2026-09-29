"""Cinematic deck helpers: gradients, glows, shadows, 3D transitions."""
import random
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC   = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14  = "http://schemas.microsoft.com/office/powerpoint/2010/main"
PNS  = "http://schemas.openxmlformats.org/presentationml/2006/main"
P_   = "p"

BG      = RGBColor(0x0B, 0x11, 0x20)
CARD    = RGBColor(0x12, 0x1D, 0x36)
ACCENT  = RGBColor(0x22, 0xD3, 0xEE)
ACCENT2 = RGBColor(0x8B, 0x5C, 0xF6)
GREEN   = RGBColor(0x34, 0xD3, 0x99)
AMBER   = RGBColor(0xFB, 0xBF, 0x24)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
TEXT    = RGBColor(0xE9, 0xF1, 0xFB)
MUTED   = RGBColor(0x8E, 0x9D, 0xB8)

FONT = "Segoe UI"

SC = {"22D3EE": ACCENT, "8B5CF6": ACCENT2, "34D399": GREEN, "FBBF24": AMBER,
      "FFFFFF": WHITE, "E9F1FB": TEXT, "F87171": RGBColor(0xF8, 0x71, 0x71),
      "D1A93C": RGBColor(0xD1, 0xA9, 0x3C), "0E74A0": RGBColor(0x0E, 0x74, 0xA0)}

prs = None


def setup():
    global prs
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _clr(hexv, alpha=None):
    if alpha is None:
        return '<a:srgbClr val="%s"/>' % hexv
    return '<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>' % (hexv, int(alpha * 1000))


def _grad(shape, gs, kind, extra):
    sp = shape._element.spPr
    for t in ("solidFill", "gradFill", "noFill"):
        e = sp.find("{%s}%s" % (A_NS, t))
        if e is not None:
            sp.remove(e)
    xml = '<a:gradFill xmlns:a="%s"><a:gsLst>%s</a:gsLst>%s</a:gradFill>' % (A_NS, gs, extra)
    el = etree.fromstring(xml)
    geom = sp.find("{%s}prstGeom" % A_NS)
    geom.addnext(el)


def inject_linear(shape, stops, angle=100):
    """stops: list of (pos, hex, alpha_or_None)."""
    gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 100000), _clr(h, a)) for p, h, a in stops)
    _grad(shape, gs, "lin", '<a:lin ang="%d" scaled="1"/>' % int(angle * 60000))


def inject_radial(shape, stops, rect=(48, 30, 52, 70)):
    gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 100000), _clr(h, a)) for p, h, a in stops)
    l, t, r, b = rect
    _grad(shape, gs, "circle",
          '<a:path path="circle"><a:fillToRect l="%d" t="%d" r="%d" b="%d"/></a:path>' % (l, t, r, b))


def inject_effects(shape, shadow=True, glow=None, shadow_alpha=55):
    sp = shape._element.spPr
    for e in sp.findall("{%s}effectLst" % A_NS):
        sp.remove(e)
    parts = []
    if shadow:
        parts.append('<a:outerShdw blurRad="130000" dist="45000" dir="5400000" '
                     'rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="%d"/>'
                     '</a:srgbClr></a:outerShdw>' % (shadow_alpha * 1000))
    if glow:
        parts.append('<a:glow rad="60000"><a:srgbClr val="%s"><a:alpha val="15000"/>'
                     '</a:srgbClr></a:glow>' % glow)
    if parts:
        sp.append(etree.fromstring('<a:effectLst xmlns:a="%s">%s</a:effectLst>' % (A_NS, "".join(parts))))


def solid_alpha(shape, alpha_pct):
    sp = shape._element.spPr
    srgb = None
    sf = sp.find("{%s}solidFill" % A_NS)
    if sf is not None:
        srgb = sf.find("{%s}srgbClr" % A_NS)
    if srgb is None:
        ln = sp.find("{%s}ln" % A_NS)
        if ln is not None:
            srgb = ln.find("{%s}srgbClr" % A_NS)
    if srgb is not None:
        etree.SubElement(srgb, "{%s}alpha" % A_NS).set("val", str(int(alpha_pct * 1000)))


def add_transition(slide, kind="cube", attr="l", spd="med", dur=650):
    tmpl = {
        "cube": ('<p14:cube dir="%s"/>', {"l", "r", "u", "d"}),
        "doors": ('<p14:doors dir="%s" orient="h"/>', {"l", "r", "u", "d"}),
        "flip": ('<p14:flip dir="%s" pres="obj"/>', {"l", "r", "u", "d"}),
        "gallery": ('<p14:gallery dir="%s"/>', {"l", "r", "u", "d"}),
        "orbit": ('<p14:orbit dir="%s"/>', {"l", "r"}),
        "ripple": ('<p14:ripple dir="%s"/>', {"center", "topLeft", "topRight",
                                              "bottomLeft", "bottomRight"}),
        "turntable": ('<p14:turntable dir="%s"/>', {"l", "r"}),
        "fracture": ("<p14:fracture/>", {None}),
    }[kind]
    inner = tmpl[0] % ("" if attr is None else attr) if "%s" in tmpl[0] else tmpl[0]
    xml = ('<mc:AlternateContent xmlns:mc="%s" xmlns:p14="%s" xmlns:%s="%s">'
           '<mc:Choice Requires="p14"><p:transition spd="%s" p14:dur="%s">%s</p:transition></mc:Choice>'
           '<mc:Fallback><p:transition spd="%s"><p:fade/></p:transition></mc:Fallback>'
           '</mc:AlternateContent>') % (MC, P14, P_, PNS, spd, dur, inner, spd)
    sld = slide._element
    clr = sld.find("{%s}clrMapOvr" % PNS)
    if clr is None:
        clr = sld.find("{%s}cSld" % PNS)
    clr.addnext(etree.fromstring(xml))


def _kills(sh):
    sh.shadow.inherit = False


def rect(slide, x, y, w, h, color, line_color=None, line_w=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color; sp.line.width = Pt(line_w or 1)
    _kills(sp)
    return sp


def card(slide, x, y, w, h, color=CARD, line_color=None, line_w=1.0,
         glow=None, shadow=True, rad=0.05):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color; sp.line.width = Pt(line_w)
    try:
        sp.adjustments[0] = rad
    except Exception:
        pass
    _kills(sp)
    inject_effects(sp, shadow=shadow, glow=glow)
    return sp


def textbox(slide, x, y, w, h, runs, size=14, bold=False, color=TEXT,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
            space_after=4, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, size, bold, color)]
    paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing; p.space_after = Pt(space_after)
        for (t, s, b, c) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(s); r.font.bold = b; r.font.color.rgb = c; r.font.name = FONT
    return tb


def grad_text(slide, x, y, w, h, txt, size, colors=("E9F1FB", "22D3EE"),
              align=PP_ALIGN.LEFT, bold=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    fl = r.font.fill
    fl.gradient()
    stops = fl.gradient_stops
    stops[0].color.rgb = SC[colors[0]]; stops[0].position = 0.0
    stops[1].color.rgb = SC[colors[1]]; stops[1].position = 1.0
    try:
        fl.gradient_angle = 90
    except Exception:
        pass
    return tb


def pill(slide, x, y, w, h, label, color=ACCENT, text_color=BG, glow=None):
    sp = card(slide, x, y, w, h, color=color, glow=glow, rad=0.5)
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(11); r.font.bold = True
    r.font.color.rgb = text_color; r.font.name = FONT
    return sp


def arrow_down(slide, x, y, color=ACCENT):
    sp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y),
                                Inches(0.28), Inches(0.24))
    sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background()
    _kills(sp); inject_effects(sp, shadow=True)
    return sp


VARIANTS = [
    (("22D3EE", 11.3, 0.6, 3.0), ("8B5CF6", 1.6, 6.9, 3.4)),
    (("8B5CF6", 0.5, 1.1, 2.6), ("22D3EE", 12.5, 6.3, 3.0)),
    (("22D3EE", 12.9, 1.4, 2.2), ("8B5CF6", 0.9, 6.6, 3.4)),
    (("22D3EE", 2.6, 0.3, 2.8), ("F59E0B", 11.4, 6.8, 2.6)),
    (("22D3EE", 7.0, 0.0, 3.6), ("8B5CF6", 12.7, 7.0, 3.0)),
]


def scene(slide, variant, seed=7):
    g = rect(slide, 0, 0, 13.333, 7.5, BG)
    inject_linear(g, [(0.0, "0D1530", None), (0.5, "0A1222", None), (1.0, "04060C", None)], angle=95)
    for (colr, cx, cy, r) in variant:
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                                   Inches(2 * r), Inches(2 * r))
        o.line.fill.background(); _kills(o)
        inject_radial(o, [(0.0, colr, 26), (0.6, colr, 7), (1.0, colr, 0)])
    rng = random.Random(seed)
    for _ in range(30):
        x = rng.uniform(0.2, 13.0); y = rng.uniform(0.2, 7.2)
        si = rng.choice([0.018, 0.026, 0.034])
        st = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(si), Inches(si))
        st.fill.solid(); st.fill.fore_color.rgb = RGBColor(0xCF, 0xE8, 0xFE)
        solid_alpha(st, rng.uniform(5, 20)); st.line.fill.background(); _kills(st)
    return g


def beam(slide, x, y, w, h, rot, alpha=7):
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = WHITE; solid_alpha(b, alpha)
    b.line.fill.background(); _kills(b); b.rotation = rot
    return b


def header(slide, eyebrow, title, num, varslide):
    scene(slide, VARIANTS[varslide % len(VARIANTS)], seed=100 + varslide)
    beam(slide, -1.2, 0.15, 15.0, 0.09, -18, alpha=6)
    textbox(slide, 0.6, 0.3, 11.5, 0.3, [("  " + eyebrow, 12, True, ACCENT)])
    grad_text(slide, 0.6, 0.55, 12.2, 0.8, title, 30, colors=("E9F1FB", "22D3EE"))
    gline = rect(slide, 0.63, 1.4, 0.7, 0.09, ACCENT)
    inject_effects(gline, shadow=False, glow="22D3EE")
    footer(slide, num)


def footer(slide, num):
    rect(slide, 0, 7.26, 13.333, 0.02, RGBColor(0x2A, 0x3A, 0x5C))
    textbox(slide, 0.6, 6.98, 9.5, 0.28,
            [("SMARTPRICE.AI  •  VEER SINGH  •  ROLL 2449476", 9, True, MUTED)])
    textbox(slide, 12.4, 6.98, 0.7, 0.28, [(str(num), 11, True, ACCENT)], align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, points, color=TEXT, size=15, gap=7, mark="  "):
    paras = []
    for pt in points:
        paras.append([(mark, size, True, ACCENT), (pt, size, False, color)])
    textbox(slide, x, y, w, h, paras, size=size, space_after=gap, line_spacing=1.18)