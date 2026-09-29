"""SmartPrice.AI — cinematic 3D-movie deck (13 slides)."""
from deck_helpers import (setup, add_transition, scene, beam, header, footer, rect, card,
                          textbox, grad_text, pill, arrow_down, bullets, inject_linear,
                          inject_effects, inject_radial, solid_alpha, BG, CARD, ACCENT,
                          ACCENT2, GREEN, AMBER, WHITE, TEXT, MUTED)
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

EYE = "BTECH CSE  •  DATA SCIENCE  •  5TH SEMESTER"
REDX = RGBColor(0xF8, 0x71, 0x71)
GLOW = {ACCENT: "22D3EE", ACCENT2: "8B5CF6", GREEN: "34D399", AMBER: "FBBF24",
        WHITE: "FFFFFF", REDX: "F87171"}
prs = setup()
BLANK = prs.slide_layouts[6]


def new_slide(kind, attr="l", spd="med", dur=650):
    s = prs.slides.add_slide(BLANK)
    add_transition(s, kind, attr, spd, dur)
    return s


def gold_line(slide, x, y, w):
    g = rect(slide, x, y, w, 0.045, ACCENT)
    inject_effects(g, shadow=False, glow="22D3EE")


# ---------------- 01 TITLE ----------------
s = new_slide("cube", "r", "med", 900)
scene(s, (( "22D3EE", 11.8, 0.7, 3.6), ("8B5CF6", 1.4, 6.8, 3.6)), seed=11)
beam(s, -2, 0.1, 17, 0.14, -16, alpha=7)
beam(s, 6.5, -1.5, 0.5, 12, -24, alpha=5)

mx, my = 9.5, 1.35
phone = card(s, mx, my, 2.6, 4.15, color=RGBColor(0x0B, 0x11, 0x20),
             line_color=RGBColor(0x3B, 0x4F, 0x78), line_w=1.5, glow="22D3EE", rad=0.09)
screen = rect(s, mx + 0.14, my + 0.24, 2.32, 3.25, RGBColor(0x16, 0x4E, 0x63))
inject_linear(screen, [(0.0, "4C1D95", None), (0.55, "164E63", None), (1.0, "0E7490", None)], angle=35)
notch = rect(s, mx + 0.92, my + 0.4, 0.76, 0.11, RGBColor(0x0A, 0x0E, 0x1A))
notch.fill.solid(); notch.fill.fore_color.rgb = RGBColor(0x0A, 0x0E, 0x1A)
notch.line.fill.background(); notch.shadow.inherit = False
w1 = rect(s, mx + 0.35, my + 0.85, 1.9, 0.32, RGBColor(0x64, 0xD8, 0xF0))
inject_linear(w1, [(0.0, "22D3EE", 70), (1.0, "22D3EE", 25)], angle=0)
w2 = rect(s, mx + 0.35, my + 1.35, 1.9, 0.10, RGBColor(0xA7, 0xF3, 0xD0)); w2.fill.solid()
w2.fill.fore_color.rgb = RGBColor(0xA7, 0xF3, 0xD0); solid_alpha(w2, 60)
w2.line.fill.background(); w2.shadow.inherit = False
w3 = rect(s, mx + 0.35, my + 1.6, 1.9, 0.10, TEXT); w3.fill.solid()
w3.fill.fore_color.rgb = TEXT; solid_alpha(w3, 40)
w3.line.fill.background(); w3.shadow.inherit = False
w4 = rect(s, mx + 0.35, my + 1.85, 1.2, 0.10, AMBER); w4.fill.solid()
w4.fill.fore_color.rgb = AMBER; solid_alpha(w4, 70)
w4.line.fill.background(); w4.shadow.inherit = False
home = rect(s, mx + 1.05, my + 3.35, 0.5, 0.05, WHITE)
home.fill.solid(); home.fill.fore_color.rgb = WHITE; solid_alpha(home, 70)
home.line.fill.background(); home.shadow.inherit = False

r1 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(mx - 0.55), Inches(my - 0.55), Inches(3.9), Inches(3.9))
r1.fill.background(); r1.line.color.rgb = ACCENT; r1.line.width = Pt(1.1)
r1.shadow.inherit = False; solid_alpha(r1, 16)
r2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(mx - 0.25), Inches(my - 0.25), Inches(3.3), Inches(3.3))
r2.fill.background(); r2.line.color.rgb = ACCENT2; r2.line.width = Pt(1.1)
r2.shadow.inherit = False; solid_alpha(r2, 14)

textbox(s, 0.9, 0.95, 7.4, 0.35, [("  B.TECH CSE (DATA SCIENCE)  •  5TH SEMESTER  •  2026", 13, True, ACCENT)])
grad_text(s, 0.9, 1.35, 8.4, 1.7, "SMARTPRICE", 60, colors=("FFFFFF", "22D3EE"))
grad_text(s, 0.9, 2.45, 8.4, 1.1, ".AI", 60, colors=("22D3EE", "8B5CF6"))
textbox(s, 0.95, 3.65, 7.6, 0.5, [("AI-Powered Price Comparison & Smart Buying Engine", 19, False, TEXT)])
textbox(s, 0.95, 4.2, 7.6, 0.4, [("Cheapest store  •  AI forecast  •  Buy Now or Wait", 13, False, MUTED)])
gold_line(s, 0.97, 4.85, 3.1)
maker = card(s, 0.9, 5.15, 7.4, 1.25, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1, glow="8B5CF6")
textbox(s, 1.25, 5.32, 6.8, 0.26, [("PROJECT  MADE  BY", 10, True, ACCENT2)])
textbox(s, 1.25, 5.62, 6.9, 0.5,
        [[("VEER SINGH", 20, True, WHITE), ("     •     Roll No. 2449476", 13, False, MUTED)]])
footer(s, 1)

# ---------------- 02 AGENDA ----------------
s = new_slide("flip", "l")
header(s, EYE, "AGENDA", 2, 1)
items = [
    ("01", "Introduction & Problem Statement"),
    ("02", "Objectives of the Project"),
    ("03", "Project Overview & Key Features"),
    ("04", "Technology Stack"),
    ("05", "The AI Engine — Linear Regression"),
    ("06", "Dataset & Pricing Engine"),
    ("07", "System Architecture"),
    ("08", "Results & Insights"),
    ("09", "Future Scope & Conclusion"),
]
for i, (num, t) in enumerate(items):
    col, row = divmod(i, 5)
    x = 0.7 + col * 6.3; y = 1.7 + row * 1.02
    card(s, x, y, 5.9, 0.82, CARD, line_color=RGBColor(0x1E, 0x2E, 0x4E), line_w=1)
    chip = card(s, x + 0.24, y + 0.16, 0.68, 0.5, color=RGBColor(0x16, 0x2A, 0x4A),
                glow="22D3EE", rad=0.3)
    tf = chip.text_frame; tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(14); r.font.bold = True
    r.font.color.rgb = ACCENT; r.font.name = "Segoe UI"
    textbox(s, x + 1.15, y + 0.15, 4.5, 0.55, [(t, 13, True, TEXT)], anchor=MSO_ANCHOR.MIDDLE)

# ---------------- 03 INTRODUCTION ----------------
s = new_slide("gallery", "r")
header(s, EYE, "INTRODUCTION & PROBLEM", 3, 2)
c1 = card(s, 0.7, 1.7, 7.5, 4.4, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
textbox(s, 1.0, 1.95, 6.9, 0.4, [("The everyday shopping problem", 17, True, ACCENT)])
bullets(s, 1.0, 2.5, 6.9, 3.3, [
    "The same smartphone is sold at many stores — brand store, Amazon, Flipkart.",
    "Prices change every week with sales, bank offers and new launches.",
    "Buyers cannot tell which store is cheapest today…",
    "…or whether the price will drop further next week.",
    "Result: people either overpay or buy too early.",
])
c2 = card(s, 8.45, 1.7, 4.15, 4.4, CARD, line_color=ACCENT, line_w=1.5, glow="22D3EE")
textbox(s, 8.75, 1.95, 3.6, 0.4, [("THE QUESTION", 15, True, ACCENT2)])
textbox(s, 8.75, 2.5, 3.6, 2.2,
        [("“Where should I buy, and WHEN should I wait for the best price?”", 16.5, True, WHITE)],
        line_spacing=1.25, space_after=0)
textbox(s, 8.75, 5.35, 3.6, 0.6, [("SmartPrice.AI answers it with data + AI.", 13, False, GREEN)])

# ---------------- 04 OBJECTIVES ----------------
s = new_slide("cube", "u")
header(s, EYE, "OBJECTIVES", 4, 3)
objs = [
    ("COMPARE", "Search any mobile and instantly see the least price across major stores.", ACCENT),
    ("FORECAST", "Predict 2-week price movement per store using from-scratch machine learning.", ACCENT2),
    ("RECOMMEND", "Give a clear verdict — BUY NOW or WAIT — with the estimated saving.", GREEN),
    ("EDUCATE", "Show the full 14-week price trend for transparent, data-backed decisions.", AMBER),
]
for i, (t, d, c) in enumerate(objs):
    col, row = divmod(i, 2)
    x = 0.7 + col * 6.35; y = 1.8 + row * 2.25
    card(s, x, y, 5.95, 1.9, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
    bar = rect(s, x, y, 0.09, 1.9, c)
    inject_effects(bar, shadow=False, glow=GLOW.get(c))
    textbox(s, x + 0.4, y + 0.3, 5.2, 0.45, [("◆  " + t, 17, True, c)])
    textbox(s, x + 0.4, y + 0.95, 5.2, 0.85, [(d, 13, False, TEXT)], line_spacing=1.2)

# ---------------- 05 OVERVIEW ----------------
s = new_slide("orbit", "r")
header(s, EYE, "PROJECT OVERVIEW", 5, 4)
textbox(s, 0.7, 1.55, 11.9, 0.7,
        [("SmartPrice.AI is a complete web app — built 100% in Python — that turns raw price data "
          "into a buying decision.", 15, False, TEXT)], line_spacing=1.2)
feats = [
    ("01", "One-search multi-store compare", "Brand store, Amazon & Flipkart — sorted cheap → costly."),
    ("02", "LOWEST price badge", "The cheapest store is flagged instantly, so you never overpay."),
    ("03", "AI price forecast", "2-week forward prediction per store, no ML libraries."),
    ("04", "BUY NOW / WAIT verdict", "Clear recommendation plus the estimated saving if you wait."),
    ("05", "14-week price history", "Transparent trend showing how each store price moved."),
    ("06", "Buy-instantly deep links", "One tap takes you to the live product page."),
]
for i, (num, t, d) in enumerate(feats):
    col, row = divmod(i, 3)
    x = 0.7 + col * 4.15; y = 2.45 + row * 2.2
    card(s, x, y, 3.85, 1.9, CARD, line_color=RGBColor(0x1E, 0x2E, 0x4E), line_w=1)
    chip = card(s, x + 0.22, y + 0.2, 0.6, 0.44, color=RGBColor(0x16, 0x2A, 0x4A),
                glow="22D3EE", rad=0.3)
    tf = chip.text_frame; tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(13); r.font.bold = True
    r.font.color.rgb = ACCENT; r.font.name = "Segoe UI"
    textbox(s, x + 0.25, y + 0.78, 3.4, 0.95, [[(t, 13, True, TEXT)], [(d, 11.5, False, MUTED)]],
            space_after=5, line_spacing=1.12)

# ---------------- 06 TECH STACK ----------------
s = new_slide("turntable", "l")
header(s, EYE, "TECHNOLOGY STACK", 6, 0)
tech = [
    ("PYTHON 3", "Single language all the way — logic, API and AI.", ACCENT),
    ("FLASK", "Lightweight web framework for the app and JSON API.", ACCENT2),
    ("HTML • CSS • JS", "Dark, modern, responsive interface.", GREEN),
    ("LINEAR REGRESSION", "From-scratch forecaster — no sklearn, no NumPy.", AMBER),
    ("GIT & GITHUB", "Version control + public code hosting.", RGBColor(0xF8, 0x71, 0x71)),
    ("RENDER", "Free cloud hosting with auto-deploy from GitHub.", WHITE),
]
for i, (t, d, c) in enumerate(tech):
    col, row = divmod(i, 3)
    x = 0.7 + col * 4.15; y = 1.85 + row * 2.3
    card(s, x, y, 3.85, 1.95, CARD, line_color=RGBColor(0x1E, 0x2E, 0x4E), line_w=1)
    bar = rect(s, x, y, 0.09, 1.95, c)
    inject_effects(bar, shadow=False, glow=GLOW.get(c))
    textbox(s, x + 0.3, y + 0.25, 3.3, 0.45, [("●  " + t, 15.5, True, c)])
    textbox(s, x + 0.3, y + 0.9, 3.3, 0.9, [(d, 12, False, TEXT)], line_spacing=1.2)
textbox(s, 0.7, 6.6, 11.9, 0.4,
        [("No external ML libraries — the whole machine-learning engine is hand-written pure Python.",
          13, True, ACCENT)], align=PP_ALIGN.CENTER)

# ---------------- 07 AI ENGINE ----------------
s = new_slide("doors", "l", "med", 800)
header(s, EYE, "THE AI ENGINE — LINEAR REGRESSION", 7, 1)
c1 = card(s, 0.7, 1.7, 6.0, 4.75, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
textbox(s, 1.0, 1.95, 5.4, 0.4, [("How the forecast works", 16, True, ACCENT)])
bullets(s, 1.0, 2.55, 5.4, 3.6, [
    "Each store has a 14-week price history.",
    "Fit the best-fit line:  y = m·x + c",
    "Slope m ₹/week and intercept c summarise the trend.",
    "Extend the line 2 weeks ahead → predicted price.",
    "Forecast vs current tells the direction: up or down.",
], size=13.5)
bar = rect(s, 7.0, 1.7, 0.06, 4.75, ACCENT2); inject_effects(bar, shadow=False, glow="8B5CF6")
c2 = card(s, 7.08, 1.7, 5.55, 4.75, CARD, line_color=ACCENT2, line_w=1.2, glow="8B5CF6")
textbox(s, 7.38, 1.95, 5.0, 0.4, [("VERDICT MATRIX", 16, True, ACCENT2)])
rows = [
    ("Slope  >  +1%", "Price RISING", "BUY NOW", GREEN),
    ("Slope  <  -1%", "Price FALLING", "WAIT", AMBER),
    ("Otherwise", "Price STABLE", "BUY NOW", GREEN),
]
yy = 2.55
for cond, sens, verdict, cv in rows:
    card(s, 7.38, yy, 4.95, 1.0, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=0.75)
    textbox(s, 7.6, yy + 0.16, 2.7, 0.6, [(cond, 13, True, WHITE)])
    textbox(s, 10.3, yy + 0.16, 1.9, 0.6, [(sens, 10.5, False, MUTED)])
    textbox(s, 7.6, yy + 0.5, 2.7, 0.35, [(verdict, 12.5, True, cv)])
    yy += 1.22
textbox(s, 7.38, 6.1, 4.95, 0.35,
        [("Evaluation: R² from residuals validates each store fit.", 11, False, MUTED)])

# ---------------- 08 DATASET ----------------
s = new_slide("ripple", "center")
header(s, EYE, "DATASET & PRICING ENGINE", 8, 2)
kpis = [("413", "Trending smartphones", ACCENT), ("1,200+", "Store prices  (3 × per phone)", ACCENT2),
        ("SEP 2026", "Prices as of — real researched listings", GREEN),
        ("14", "Weeks of price history per store", AMBER)]
for i, (n, d, c) in enumerate(kpis):
    x = 0.7 + i * 3.1
    card(s, x, 1.7, 2.85, 1.75, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
    textbox(s, x + 0.2, 1.98, 2.45, 0.6, [(n, 28, True, c)], align=PP_ALIGN.CENTER)
    textbox(s, x + 0.2, 2.7, 2.45, 0.6, [(d, 11, False, TEXT)], align=PP_ALIGN.CENTER)
c1 = card(s, 0.7, 3.75, 11.9, 2.85, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
textbox(s, 1.0, 4.0, 11.2, 0.4, [("How prices are sourced", 16, True, ACCENT)])
bullets(s, 1.0, 4.55, 11.2, 2.0, [
    "Current prices are researched from real live listings — brand stores, Amazon India & Flipkart "
    "(price trackers: 91mobiles, Smartprix, Gadgets360).",
    "Each phone keeps a 14-week history; simulated but anchored so the LAST week equals the real current price.",
    "Every seller deep-links to its live product page → the comparison stays actionable, not theoretical.",
], size=13)
textbox(s, 0.7, 6.72, 11.9, 0.35,
        [("Data honesty: prices are labelled “as of September 2026” right inside the UI.", 12, True, AMBER)],
        align=PP_ALIGN.CENTER)

# ---------------- 09 ARCHITECTURE ----------------
s = new_slide("gallery", "l")
header(s, EYE, "SYSTEM ARCHITECTURE", 9, 3)
layers = [
    ("1 — PRESENTATION LAYER", "Browser UI  •  HTML / CSS / JavaScript  •  dark interactive cards", ACCENT),
    ("2 — APPLICATION LAYER", "Flask app  •  routes  •  JSON API  (/api/search, /api/catalog)", ACCENT2),
    ("3 — BUSINESS LOGIC", "search_products  •  enrich  •  BUY NOW / WAIT verdict  •  savings", GREEN),
    ("4 — AI ENGINE", "Linear regression fit  •  2-week forecast  •  R² validation", AMBER),
    ("5 — DATA LAYER", "Dataset  •  413 phones  •  1,200+ prices  •  seeded history", RGBColor(0xF8, 0x71, 0x71)),
]
y = 1.6
for i, (t, d, ac) in enumerate(layers):
    card(s, 1.7, y, 9.9, 0.78, BG, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
    bar = rect(s, 1.7, y, 0.09, 0.78, ac)
    inject_effects(bar, shadow=False, glow=GLOW.get(ac))
    textbox(s, 2.0, y + 0.1, 9.2, 0.3, [(t, 13.5, True, WHITE)])
    textbox(s, 2.0, y + 0.4, 9.2, 0.3, [(d, 11, False, MUTED)])
    if i < len(layers) - 1:
        arrow_down(s, 6.6, y + 0.82, ACCENT)
    y += 1.12
textbox(s, 0.7, 7.05, 11.9, 0.3,
        [("Clear separation → easy to extend, test and redeploy automatically on Render.", 11.5, False, MUTED)],
        align=PP_ALIGN.CENTER)

# ---------------- 10 RESULTS ----------------
s = new_slide("flip", "r")
header(s, EYE, "RESULTS & INSIGHTS", 10, 4)
c1 = card(s, 0.7, 1.7, 7.6, 4.75, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1)
textbox(s, 1.0, 1.95, 7.0, 0.4, [("LIVE RESULT — search “iPhone 16 Pro”", 16, True, ACCENT)])
sample = [("Flipkart", "₹ 1,05,900", "★ LOWEST", GREEN), ("Amazon India", "₹ 1,12,900", "", None),
          ("Apple Store", "₹ 1,19,900", "", None)]
yy = 2.55
for name, price, tag, cv in sample:
    card(s, 1.0, yy, 7.0, 0.88, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=0.75)
    textbox(s, 1.25, yy + 0.19, 3.3, 0.5, [(name, 14, True, WHITE)])
    textbox(s, 4.4, yy + 0.19, 1.9, 0.5, [(price, 14.5, True, ACCENT)])
    if tag:
        pill(s, 6.55, yy + 0.2, 1.3, 0.46, tag, color=cv)
    yy += 1.08
c2 = card(s, 8.5, 1.7, 4.1, 4.75, CARD, line_color=ACCENT, line_w=1.5, glow="22D3EE")
textbox(s, 8.8, 1.95, 3.5, 0.4, [("AI VERDICTS IN ACTION", 16, True, ACCENT)])
textbox(s, 8.8, 2.55, 3.5, 3.6, [
    [("📈 Rising trend →", 14, True, GREEN), ("  buy now before it gets costlier.", 12.5, False, TEXT)],
    [("📉 Falling trend →", 14, True, AMBER), ("  wait ~2 weeks, keep the shown saving.", 12.5, False, TEXT)],
    [("➖ Stable trend →", 14, True, GREEN), ("  already a good time to buy.", 12.5, False, TEXT)],
], line_spacing=1.25, space_after=18)
textbox(s, 0.7, 6.65, 11.9, 0.35,
        [("Every result carries a 2-week price forecast for each store, plus a savings figure when you wait.",
          12, False, MUTED)], align=PP_ALIGN.CENTER)

# ---------------- 11 FUTURE ----------------
s = new_slide("orbit", "l")
header(s, EYE, "FUTURE SCOPE", 11, 0)
future = [
    ("Live price scraping", "Refresh all prices automatically every few hours instead of researched snapshots."),
    ("More categories", "Laptops, TVs, smartwatches, accessories, appliances."),
    ("Price-drop alerts", "Email / notification when a watched product hits its target price."),
    ("Real-history charts", "Train on true daily histories collected over time."),
    ("Android app", "Native app + browser extension for one-tap comparisons."),
]
yy = 1.7
for t, d in future:
    card(s, 0.7, yy, 11.9, 0.82, CARD, line_color=RGBColor(0x1E, 0x2E, 0x4E), line_w=1)
    textbox(s, 1.0, yy + 0.14, 4.6, 0.55, [("→  " + t, 15, True, ACCENT)])
    textbox(s, 5.6, yy + 0.17, 6.9, 0.5, [(d, 12, False, TEXT)])
    yy += 1.02

# ---------------- 12 CONCLUSION ----------------
s = new_slide("doors", "r", "med", 800)
header(s, EYE, "CONCLUSION", 12, 1)
c1 = card(s, 1.0, 1.8, 11.3, 4.45, CARD, line_color=ACCENT, line_w=1.5, glow="22D3EE")
textbox(s, 1.4, 2.15, 10.5, 0.45, [("WHAT THIS PROJECT DEMONSTRATES", 17, True, ACCENT)])
bullets(s, 1.4, 2.8, 10.5, 3.3, [
    "A complete, end-to-end AI product built in pure Python — from data collection to deployed web app.",
    "Applied data-science skills: data research, cleaning, modelling, forecasting and R² evaluation.",
    "The from-scratch linear-regression engine shows the real maths behind ML — no black-box libraries.",
    "It answers a practical question: where to buy a phone, and when — for the best price.",
    "Runs on a free stack (GitHub + Render) and auto-deploys on every commit.",
], size=14, gap=9)

# ---------------- 13 THANK YOU ----------------
s = new_slide("fracture", None, "med", 900)
scene(s, (( "22D3EE", 6.7, 1.2, 3.6), ("8B5CF6", 11.5, 6.3, 3.4)), seed=23)
beam(s, -2, 0.1, 17, 0.14, -16, alpha=7)
beam(s, 8.5, -1.5, 0.5, 12, 20, alpha=5)
tx = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.2), Inches(2.2), Inches(5.0), Inches(5.0))
tx.fill.background(); tx.line.color.rgb = ACCENT; tx.line.width = Pt(1.2)
tx.shadow.inherit = False; solid_alpha(tx, 14)
grad_text(s, 0.8, 1.5, 11.5, 1.3, "THANK YOU", 58, colors=("FFFFFF", "22D3EE"), align=PP_ALIGN.CENTER)
textbox(s, 0.8, 3.15, 11.5, 0.5,
        [("SmartPrice.AI — AI-Powered Price Comparison & Smart Buying Engine", 18, False, ACCENT)],
        align=PP_ALIGN.CENTER)
gold_line(s, 5.8, 4.0, 1.7)
maker = card(s, 3.4, 4.5, 6.5, 1.35, CARD, line_color=RGBColor(0x2A, 0x3A, 0x5C), line_w=1, glow="8B5CF6")
textbox(s, 3.7, 4.68, 5.9, 0.3, [("PROJECT  MADE  BY", 10, True, ACCENT2)], align=PP_ALIGN.CENTER)
textbox(s, 3.7, 4.98, 5.9, 0.5, [("VEER SINGH", 21, True, WHITE)], align=PP_ALIGN.CENTER)
textbox(s, 3.7, 5.5, 5.9, 0.35,
        [("B.Tech CSE — Data Science  •  5th Semester  •  Roll No. 2449476", 12, False, MUTED)],
        align=PP_ALIGN.CENTER)
textbox(s, 0.8, 6.35, 11.5, 0.4, [("Questions & suggestions are warmly welcome.", 13, False, MUTED)],
        align=PP_ALIGN.CENTER)
footer(s, 13)


cp = prs.core_properties
cp.title = "SmartPrice.AI — Project Presentation"
cp.author = "Veer Singh"
cp.subject = "AI-Powered Price Comparison & Smart Buying Engine"

OUT = "SmartPrice_AI_Presentation.pptx"
prs.save(OUT)
print("Saved", OUT, "with", len(prs.slides._sldIdLst), "slides")